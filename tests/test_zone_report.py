"""
Zone Report Tests
=================

Tests for zone_report.py and zone_excel.py modules including:
- ZoneLccaReport generation
- Category and bedroom summaries
- CUAC allowance reports
- Text formatting
- Excel-ready data export
- Excel workbook generation

Run with:
    pytest tests/test_zone_report.py -v
"""

import pytest
from typing import List
import tempfile
from pathlib import Path


# =============================================================================
# TEST FIXTURES
# =============================================================================

def create_sample_zones():
    """Create sample zones for testing."""
    from eco_tools.lcca.zone_energy import ZoneEnergySummary
    from eco_tools.lcca.cuac.models import ZoneType
    from eco_tools.lcca.res_other.models import CommonAreaCategory

    return [
        # Dwelling units
        ZoneEnergySummary(
            zone_name="Unit 101 (1BR)",
            zone_type=ZoneType.DWELLING_UNIT,
            elec_kwh=4000.0,
            gas_therm=150.0,
            area_sqft=600.0,
            num_bedrooms=1,
        ),
        ZoneEnergySummary(
            zone_name="Unit 102 (1BR)",
            zone_type=ZoneType.DWELLING_UNIT,
            elec_kwh=4200.0,
            gas_therm=160.0,
            area_sqft=620.0,
            num_bedrooms=1,
        ),
        ZoneEnergySummary(
            zone_name="Unit 201 (2BR)",
            zone_type=ZoneType.DWELLING_UNIT,
            elec_kwh=5500.0,
            gas_therm=200.0,
            area_sqft=850.0,
            num_bedrooms=2,
        ),
        ZoneEnergySummary(
            zone_name="Unit 202 (2BR)",
            zone_type=ZoneType.DWELLING_UNIT,
            elec_kwh=5300.0,
            gas_therm=190.0,
            area_sqft=820.0,
            num_bedrooms=2,
        ),
        # Common areas
        ZoneEnergySummary(
            zone_name="Lobby",
            zone_type=ZoneType.COMMON_AREA,
            category=CommonAreaCategory.LOBBY,
            elec_kwh=2000.0,
            gas_therm=0.0,
            area_sqft=400.0,
        ),
        ZoneEnergySummary(
            zone_name="Corridor L1",
            zone_type=ZoneType.COMMON_AREA,
            category=CommonAreaCategory.CORRIDOR,
            elec_kwh=1200.0,
            gas_therm=0.0,
            area_sqft=300.0,
        ),
        ZoneEnergySummary(
            zone_name="Corridor L2",
            zone_type=ZoneType.COMMON_AREA,
            category=CommonAreaCategory.CORRIDOR,
            elec_kwh=1300.0,
            gas_therm=0.0,
            area_sqft=320.0,
        ),
        ZoneEnergySummary(
            zone_name="Fitness Center",
            zone_type=ZoneType.COMMON_AREA,
            category=CommonAreaCategory.FITNESS,
            elec_kwh=3500.0,
            gas_therm=0.0,
            area_sqft=600.0,
        ),
    ]


def create_sample_tariff():
    """Create sample tariff for testing."""
    from eco_tools.lcca.tariffs import create_pge_e_tou_c
    return create_pge_e_tou_c()


def create_sample_cuac_config():
    """Create sample CUAC config for testing."""
    from eco_tools.lcca.cuac.models import CuacConfig
    return CuacConfig(
        elec_utility="Pacific Gas and Electric Company",
        elec_tariff="E-TOU-C",
        affordable_pv_dc_sys_size=20.0,
        pct_indiv_unit_pv_by_bedrms={1: 4.0, 2: 6.0},
        pv_billing_option="Virtual Net Energy Metering",
    )


# =============================================================================
# ZONE LCCA REPORT TESTS
# =============================================================================

class TestZoneLccaReport:
    """Tests for ZoneLccaReport data class."""

    def test_basic_construction(self):
        """Test basic report construction."""
        from eco_tools.lcca.zone_report import ZoneLccaReport

        report = ZoneLccaReport(
            building_name="Test Building",
            building_type="Multifamily",
            total_area_sqft=5000.0,
            dwelling_unit_count=10,
            common_area_count=5,
        )

        assert report.building_name == "Test Building"
        assert report.building_type == "Multifamily"
        assert report.dwelling_unit_count == 10

    def test_eui_calculation(self):
        """Test EUI auto-calculation."""
        from eco_tools.lcca.zone_report import ZoneLccaReport

        report = ZoneLccaReport(
            total_area_sqft=1000.0,
            total_elec_kwh=1000.0,  # 3412 kBtu
            total_gas_therm=100.0,  # 10000 kBtu
        )

        # EUI = (3412 + 10000) / 1000 = 13.412
        assert abs(report.total_eui_kbtu_sqft - 13.412) < 0.01

    def test_has_lcca_property(self):
        """Test has_lcca property."""
        from eco_tools.lcca.zone_report import ZoneLccaReport
        from eco_tools.lcca.calculators import LccaResults

        report_no_lcca = ZoneLccaReport()
        assert report_no_lcca.has_lcca is False

        report_with_lcca = ZoneLccaReport(
            lcca_results=LccaResults(npv=100000.0)
        )
        assert report_with_lcca.has_lcca is True

    def test_has_cuac_property(self):
        """Test has_cuac property."""
        from eco_tools.lcca.zone_report import ZoneLccaReport, CuacSummaryReport

        report_no_cuac = ZoneLccaReport()
        assert report_no_cuac.has_cuac is False

        report_with_cuac = ZoneLccaReport(
            cuac_summary=CuacSummaryReport()
        )
        assert report_with_cuac.has_cuac is True

    def test_get_executive_summary(self):
        """Test executive summary generation."""
        from eco_tools.lcca.zone_report import ZoneLccaReport

        report = ZoneLccaReport(
            building_name="Test Building",
            total_area_sqft=5000.0,
            dwelling_unit_count=10,
            total_annual_cost=50000.0,
        )

        summary = report.get_executive_summary()

        assert summary['building_name'] == "Test Building"
        assert summary['total_area_sqft'] == 5000.0
        assert summary['total_annual_cost'] == 50000.0

    def test_get_cost_breakdown(self):
        """Test cost breakdown generation."""
        from eco_tools.lcca.zone_report import ZoneLccaReport

        report = ZoneLccaReport(
            dwelling_unit_cost=40000.0,
            common_area_cost=10000.0,
            total_annual_cost=50000.0,
        )

        breakdown = report.get_cost_breakdown()

        assert breakdown['dwelling_units'] == 40000.0
        assert breakdown['common_areas'] == 10000.0
        assert breakdown['total'] == 50000.0

    def test_to_dict(self):
        """Test dictionary serialization."""
        from eco_tools.lcca.zone_report import ZoneLccaReport

        report = ZoneLccaReport(
            building_name="Test Building",
            total_annual_cost=50000.0,
        )

        d = report.to_dict()

        assert 'metadata' in d
        assert 'building' in d
        assert 'summary' in d
        assert d['building']['name'] == "Test Building"


# =============================================================================
# REPORT GENERATION TESTS
# =============================================================================

class TestReportGeneration:
    """Tests for report generation functions."""

    def test_generate_zone_lcca_report(self):
        """Test full report generation."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Test Apartments",
            building_type="Multifamily",
            climate_zone="CZ12",
        )

        assert report.building_name == "Test Apartments"
        assert report.dwelling_unit_count == 4
        assert report.common_area_count == 4
        assert report.total_annual_cost > 0

    def test_category_summaries_generated(self):
        """Test category summaries are generated."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        assert len(report.category_summaries) == 2  # Dwelling + Common
        assert report.category_summaries[0].category_name == "Dwelling Units"
        assert report.category_summaries[1].category_name == "Common Areas"

    def test_bedroom_summaries_generated(self):
        """Test bedroom summaries are generated."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        assert len(report.bedroom_summaries) == 2  # 1BR + 2BR
        assert report.bedroom_summaries[0].num_bedrooms == 1
        assert report.bedroom_summaries[0].unit_count == 2
        assert report.bedroom_summaries[1].num_bedrooms == 2
        assert report.bedroom_summaries[1].unit_count == 2

    def test_common_area_summaries_generated(self):
        """Test common area summaries are generated."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report
        from eco_tools.lcca.res_other.models import CommonAreaCategory

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        # Should have LOBBY, CORRIDOR, FITNESS
        assert len(report.common_area_summaries) == 3

        categories = {cas.category for cas in report.common_area_summaries}
        assert CommonAreaCategory.LOBBY in categories
        assert CommonAreaCategory.CORRIDOR in categories
        assert CommonAreaCategory.FITNESS in categories

    def test_cuac_summary_generated(self):
        """Test CUAC summary is generated when config provided."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            cuac_config=cuac_config,
        )

        assert report.has_cuac
        assert report.cuac_summary.total_pv_kwdc == 20.0
        assert 1 in report.cuac_summary.allowances_by_bedroom
        assert 2 in report.cuac_summary.allowances_by_bedroom


# =============================================================================
# CUAC SUMMARY TESTS
# =============================================================================

class TestCuacSummary:
    """Tests for CUAC summary functionality."""

    def test_cuac_summary_construction(self):
        """Test CUAC summary construction."""
        from eco_tools.lcca.zone_report import CuacSummaryReport

        summary = CuacSummaryReport(
            utility_name="PG&E",
            tariff_name="E-TOU-C",
            total_pv_kwdc=20.0,
        )

        assert summary.utility_name == "PG&E"
        assert summary.total_pv_kwdc == 20.0

    def test_get_allowance_table(self):
        """Test allowance table generation."""
        from eco_tools.lcca.zone_report import CuacSummaryReport
        from eco_tools.lcca.zone_energy import CuacAllowanceResult

        summary = CuacSummaryReport()
        summary.allowances_by_bedroom = {
            1: CuacAllowanceResult(
                num_bedrooms=1,
                elec_cooling_allowance=30.0,
                elec_other_allowance=20.0,
                gas_heating_allowance=25.0,
                pv_credit_monthly=10.0,
            ),
            2: CuacAllowanceResult(
                num_bedrooms=2,
                elec_cooling_allowance=40.0,
                elec_other_allowance=30.0,
                gas_heating_allowance=35.0,
                pv_credit_monthly=15.0,
            ),
        }

        table = summary.get_allowance_table()

        assert len(table) == 2
        assert table[0]['bedrooms'] == 1
        assert table[0]['gross_total'] == 75.0  # 30+20+25
        assert table[0]['pv_credit'] == 10.0
        assert table[0]['net_allowance'] == 65.0  # 75-10


# =============================================================================
# TEXT FORMATTING TESTS
# =============================================================================

class TestTextFormatting:
    """Tests for text formatting functions."""

    def test_format_zone_report(self):
        """Test zone report text formatting."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report, format_zone_report

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Test Apartments",
        )

        formatted = format_zone_report(report)

        assert "Zone-Level LCCA Report" in formatted
        assert "Test Apartments" in formatted
        assert "EXECUTIVE SUMMARY" in formatted
        assert "COST BY CATEGORY" in formatted
        assert "Dwelling Units" in formatted

    def test_format_cuac_allowance_table(self):
        """Test CUAC allowance table formatting."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report, format_cuac_allowance_table
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            cuac_config=cuac_config,
        )

        formatted = format_cuac_allowance_table(report)

        assert "CUAC UTILITY ALLOWANCE SCHEDULE" in formatted
        assert "1" in formatted  # 1-BR
        assert "2" in formatted  # 2-BR

    def test_format_zone_detail_table(self):
        """Test zone detail table formatting."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report, format_zone_detail_table
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        formatted = format_zone_detail_table(report)

        assert "ZONE DETAIL TABLE" in formatted
        assert "Unit 101" in formatted
        assert "Lobby" in formatted
        assert "TOTAL" in formatted


# =============================================================================
# EXCEL DATA EXPORT TESTS
# =============================================================================

class TestExcelDataExport:
    """Tests for Excel-ready data export."""

    def test_get_excel_sheets_data(self):
        """Test Excel sheets data generation."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report, get_excel_sheets_data
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            cuac_config=cuac_config,
        )

        sheets = get_excel_sheets_data(report)

        assert 'Summary' in sheets
        assert 'Zone Detail' in sheets
        assert 'Dwelling Units' in sheets
        assert 'Common Areas' in sheets
        assert 'CUAC Allowances' in sheets

    def test_summary_sheet_data(self):
        """Test summary sheet data structure."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report, get_excel_sheets_data
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Test Building",
        )

        sheets = get_excel_sheets_data(report)
        summary = sheets['Summary']

        assert len(summary) > 0
        metrics = {row['Metric']: row['Value'] for row in summary}
        assert metrics['Building Name'] == "Test Building"
        assert 'Total Area (sqft)' in metrics

    def test_zone_detail_sheet_data(self):
        """Test zone detail sheet data structure."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report, get_excel_sheets_data
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        sheets = get_excel_sheets_data(report)
        zone_detail = sheets['Zone Detail']

        assert len(zone_detail) == 8  # 4 dwelling + 4 common
        assert 'Zone Name' in zone_detail[0]
        assert 'Net Annual Cost ($)' in zone_detail[0]


# =============================================================================
# EXCEL WORKBOOK EXPORT TESTS
# =============================================================================

class TestExcelWorkbookExport:
    """Tests for Excel workbook export."""

    def test_export_to_excel_creates_file(self):
        """Test Excel export creates file."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report
        from eco_tools.lcca.zone_excel import export_zone_report_to_excel

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Test Building",
        )

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            output_path = f.name

        try:
            result_path = export_zone_report_to_excel(report, output_path)
            assert result_path.exists()
            assert result_path.suffix == '.xlsx'
        finally:
            Path(output_path).unlink(missing_ok=True)

    def test_export_with_cuac(self):
        """Test Excel export includes CUAC sheet."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report
        from eco_tools.lcca.zone_excel import export_zone_report_to_excel

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            cuac_config=cuac_config,
        )

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            output_path = f.name

        try:
            result_path = export_zone_report_to_excel(report, output_path)
            assert result_path.exists()

            # Verify workbook has CUAC sheet
            from openpyxl import load_workbook
            wb = load_workbook(result_path)
            assert "CUAC Allowances" in wb.sheetnames
            wb.close()
        finally:
            Path(output_path).unlink(missing_ok=True)

    def test_export_without_charts(self):
        """Test Excel export without charts."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report
        from eco_tools.lcca.zone_excel import export_zone_report_to_excel

        zones = create_sample_zones()
        tariff = create_sample_tariff()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
        )

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            output_path = f.name

        try:
            result_path = export_zone_report_to_excel(
                report, output_path, include_charts=False
            )
            assert result_path.exists()
        finally:
            Path(output_path).unlink(missing_ok=True)


# =============================================================================
# CONVENIENCE FUNCTION TESTS
# =============================================================================

class TestConvenienceFunctions:
    """Tests for convenience functions."""

    def test_create_zone_report_from_allocations(self):
        """Test creating report from allocations."""
        from eco_tools.lcca.zone_report import create_zone_report_from_allocations
        from eco_tools.lcca.cuac.models import DwellUnitAllocation

        allocations = [
            DwellUnitAllocation(
                zone_name="Unit 1",
                conditioning_type="Conditioned",
                space_function="High-Rise Residential Living Spaces",
                pv_batt_bldg_type="Highrise Multifamily",
                floor_area_sqft=800.0,
            ),
            DwellUnitAllocation(
                zone_name="Unit 2",
                conditioning_type="Conditioned",
                space_function="High-Rise Residential Living Spaces",
                pv_batt_bldg_type="Highrise Multifamily",
                floor_area_sqft=750.0,
            ),
            DwellUnitAllocation(
                zone_name="Lobby",
                conditioning_type="Conditioned",
                space_function="Lobby - Main Entry",
                pv_batt_bldg_type="Highrise Multifamily",
                floor_area_sqft=400.0,
            ),
        ]

        tariff = create_sample_tariff()

        report = create_zone_report_from_allocations(
            zone_allocations=allocations,
            building_elec_kwh=20000.0,
            building_gas_therm=1000.0,
            tariff=tariff,
            building_name="Test Building",
        )

        assert report.building_name == "Test Building"
        assert len(report.zone_results) == 3
        assert report.total_annual_cost > 0


# =============================================================================
# INTEGRATION TESTS
# =============================================================================

class TestReportIntegration:
    """Integration tests for zone reporting."""

    def test_full_workflow_text_report(self):
        """Test complete workflow with text report."""
        from eco_tools.lcca.zone_report import (
            generate_zone_lcca_report,
            format_zone_report,
            format_zone_detail_table,
        )

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        # Generate report
        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Sunset Apartments",
            building_type="Multifamily",
            climate_zone="CZ12",
            cuac_config=cuac_config,
            prepared_by="ECO Tools",
        )

        # Verify structure
        assert report.dwelling_unit_count == 4
        assert report.common_area_count == 4
        assert len(report.bedroom_summaries) == 2
        assert len(report.common_area_summaries) == 3

        # Generate text reports
        main_report = format_zone_report(report)
        detail_table = format_zone_detail_table(report)

        # Verify content
        assert "Sunset Apartments" in main_report
        assert "1-Bedroom" in main_report
        assert "CUAC" in main_report
        assert "Unit 101" in detail_table

    def test_full_workflow_excel_export(self):
        """Test complete workflow with Excel export."""
        from eco_tools.lcca.zone_report import generate_zone_lcca_report
        from eco_tools.lcca.zone_excel import export_zone_report_to_excel
        from openpyxl import load_workbook

        zones = create_sample_zones()
        tariff = create_sample_tariff()
        cuac_config = create_sample_cuac_config()

        report = generate_zone_lcca_report(
            zone_summaries=zones,
            tariff=tariff,
            building_name="Sunset Apartments",
            cuac_config=cuac_config,
        )

        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            output_path = f.name

        try:
            result_path = export_zone_report_to_excel(report, output_path)

            # Load and verify workbook
            wb = load_workbook(result_path)

            # Check sheets exist
            assert "Summary" in wb.sheetnames
            assert "Zone Detail" in wb.sheetnames
            assert "Dwelling Units" in wb.sheetnames
            assert "Common Areas" in wb.sheetnames
            assert "CUAC Allowances" in wb.sheetnames

            # Check Summary sheet content
            ws = wb["Summary"]
            assert "Sunset Apartments" in str(ws['B3'].value or ws['B4'].value or ws['B5'].value)

            wb.close()
        finally:
            Path(output_path).unlink(missing_ok=True)


if __name__ == '__main__':
    pytest.main([__file__, '-v'])
