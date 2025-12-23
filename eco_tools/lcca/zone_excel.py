"""
Zone-Level LCCA Excel Export
============================

Multi-sheet Excel workbook export for zone-level LCCA reports including:
- Summary dashboard
- Zone detail table
- Dwelling unit breakdown by bedroom
- Common area breakdown by category
- CUAC utility allowance schedule
- Charts and visualizations
"""

from __future__ import annotations
from pathlib import Path
from typing import Dict, List, Optional, Any, TYPE_CHECKING

try:
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, Alignment, Border, Side, PatternFill, NamedStyle
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.series import DataPoint
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

if TYPE_CHECKING:
    from .zone_report import ZoneLccaReport


# =============================================================================
# STYLE DEFINITIONS
# =============================================================================

# Color scheme
COLORS = {
    'primary': '4472C4',      # Blue
    'secondary': '5B9BD5',    # Light blue
    'accent1': '70AD47',      # Green
    'accent2': 'FFC000',      # Yellow/Gold
    'accent3': 'ED7D31',      # Orange
    'light_gray': 'F2F2F2',
    'dark_gray': '404040',
    'white': 'FFFFFF',
    'dwelling': '4472C4',     # Blue for dwelling units
    'common': '70AD47',       # Green for common areas
}


def _check_openpyxl():
    """Check if openpyxl is available."""
    if not OPENPYXL_AVAILABLE:
        raise ImportError(
            "openpyxl is required for Excel export. "
            "Install with: pip install openpyxl"
        )


def _create_styles(wb: "Workbook") -> Dict[str, Any]:
    """Create named styles for the workbook."""
    styles = {}

    # Title style
    title_style = NamedStyle(name="zone_title")
    title_style.font = Font(bold=True, size=16, color=COLORS['dark_gray'])
    title_style.alignment = Alignment(horizontal="left", vertical="center")
    wb.add_named_style(title_style)
    styles["title"] = title_style

    # Section header style
    section_style = NamedStyle(name="zone_section")
    section_style.font = Font(bold=True, size=12, color=COLORS['white'])
    section_style.fill = PatternFill(
        start_color=COLORS['primary'],
        end_color=COLORS['primary'],
        fill_type="solid"
    )
    section_style.alignment = Alignment(horizontal="left", vertical="center")
    wb.add_named_style(section_style)
    styles["section"] = section_style

    # Table header style
    header_style = NamedStyle(name="zone_header")
    header_style.font = Font(bold=True, size=10, color=COLORS['white'])
    header_style.fill = PatternFill(
        start_color=COLORS['secondary'],
        end_color=COLORS['secondary'],
        fill_type="solid"
    )
    header_style.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_style.border = Border(
        bottom=Side(style="thin", color=COLORS['dark_gray'])
    )
    wb.add_named_style(header_style)
    styles["header"] = header_style

    # Currency style
    currency_style = NamedStyle(name="zone_currency")
    currency_style.number_format = '"$"#,##0.00'
    currency_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(currency_style)
    styles["currency"] = currency_style

    # Currency (no cents) style
    currency_int_style = NamedStyle(name="zone_currency_int")
    currency_int_style.number_format = '"$"#,##0'
    currency_int_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(currency_int_style)
    styles["currency_int"] = currency_int_style

    # Number style
    number_style = NamedStyle(name="zone_number")
    number_style.number_format = "#,##0"
    number_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(number_style)
    styles["number"] = number_style

    # Decimal style
    decimal_style = NamedStyle(name="zone_decimal")
    decimal_style.number_format = "#,##0.0"
    decimal_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(decimal_style)
    styles["decimal"] = decimal_style

    # Percent style
    percent_style = NamedStyle(name="zone_percent")
    percent_style.number_format = "0.0%"
    percent_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(percent_style)
    styles["percent"] = percent_style

    # Alternating row style
    alt_row_style = NamedStyle(name="zone_alt_row")
    alt_row_style.fill = PatternFill(
        start_color=COLORS['light_gray'],
        end_color=COLORS['light_gray'],
        fill_type="solid"
    )
    wb.add_named_style(alt_row_style)
    styles["alt_row"] = alt_row_style

    return styles


# =============================================================================
# WORKBOOK CREATION
# =============================================================================

def export_zone_report_to_excel(
    report: "ZoneLccaReport",
    output_path: str,
    include_charts: bool = True,
) -> Path:
    """
    Export zone LCCA report to Excel workbook.

    Args:
        report: ZoneLccaReport to export
        output_path: Output file path (.xlsx)
        include_charts: Include charts in workbook

    Returns:
        Path to created file
    """
    _check_openpyxl()

    wb = Workbook()
    styles = _create_styles(wb)

    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Create sheets
    _create_summary_sheet(wb, report, styles)
    _create_zone_detail_sheet(wb, report, styles)
    _create_dwelling_units_sheet(wb, report, styles, include_charts)
    _create_common_areas_sheet(wb, report, styles, include_charts)

    if report.has_cuac:
        _create_cuac_sheet(wb, report, styles)

    # Save workbook
    output_path = Path(output_path)
    wb.save(output_path)

    return output_path


def _create_summary_sheet(
    wb: "Workbook",
    report: "ZoneLccaReport",
    styles: Dict[str, Any]
):
    """Create summary dashboard sheet."""
    ws = wb.create_sheet("Summary")

    row = 1

    # Title
    ws.merge_cells(f'A{row}:F{row}')
    ws[f'A{row}'] = report.report_title
    ws[f'A{row}'].style = styles["title"]
    row += 2

    # Building Info Section
    ws[f'A{row}'] = "BUILDING INFORMATION"
    ws[f'A{row}'].style = styles["section"]
    ws.merge_cells(f'A{row}:F{row}')
    row += 1

    info_rows = [
        ("Building Name", report.building_name or "N/A"),
        ("Building Type", report.building_type),
        ("Climate Zone", report.climate_zone or "N/A"),
        ("Report Date", report.report_date),
        ("Utility", report.utility_name),
        ("Tariff", report.tariff_name),
    ]
    for label, value in info_rows:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

    row += 1

    # Building Metrics Section
    ws[f'A{row}'] = "BUILDING METRICS"
    ws[f'A{row}'].style = styles["section"]
    ws.merge_cells(f'A{row}:F{row}')
    row += 1

    metrics = [
        ("Total Conditioned Area", f"{report.total_area_sqft:,.0f}", "sqft"),
        ("Dwelling Units", str(report.dwelling_unit_count), ""),
        ("Common Area Zones", str(report.common_area_count), ""),
        ("Annual Electricity", f"{report.total_elec_kwh:,.0f}", "kWh"),
        ("Annual Gas", f"{report.total_gas_therm:,.0f}", "therms"),
        ("Site EUI", f"{report.total_eui_kbtu_sqft:.1f}", "kBtu/sqft"),
    ]
    for label, value, unit in metrics:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'C{row}'] = unit
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

    row += 1

    # Annual Costs Section
    ws[f'A{row}'] = "ANNUAL ENERGY COSTS"
    ws[f'A{row}'].style = styles["section"]
    ws.merge_cells(f'A{row}:F{row}')
    row += 1

    cost_rows = [
        ("Dwelling Units", report.dwelling_unit_cost),
        ("Common Areas", report.common_area_cost),
        ("Total Annual Cost", report.total_annual_cost),
    ]
    for label, value in cost_rows:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = value
        ws[f'B{row}'].style = styles["currency"]
        ws[f'A{row}'].font = Font(bold=True)
        if label == "Total Annual Cost":
            ws[f'A{row}'].font = Font(bold=True, size=11)
            ws[f'B{row}'].font = Font(bold=True, size=11)
        row += 1

    # LCCA Metrics (if available)
    if report.has_lcca:
        row += 1
        ws[f'A{row}'] = "LIFECYCLE ANALYSIS"
        ws[f'A{row}'].style = styles["section"]
        ws.merge_cells(f'A{row}:F{row}')
        row += 1

        lcca = report.lcca_results
        lcca_rows = [
            ("Net Present Value (NPV)", f"${lcca.npv:,.0f}"),
            ("Internal Rate of Return", f"{(lcca.irr or 0) * 100:.1f}%" if lcca.irr else "N/A"),
            ("Simple Payback", f"{lcca.simple_payback_years:.1f} years" if lcca.simple_payback_years else "N/A"),
            ("Lifecycle Savings", f"${lcca.lifecycle_savings:,.0f}"),
        ]
        for label, value in lcca_rows:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            ws[f'A{row}'].font = Font(bold=True)
            row += 1

    # Adjust column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 12


def _create_zone_detail_sheet(
    wb: "Workbook",
    report: "ZoneLccaReport",
    styles: Dict[str, Any]
):
    """Create zone detail table sheet."""
    ws = wb.create_sheet("Zone Detail")

    # Headers
    headers = [
        "Zone Name", "Type", "Category", "Area (sqft)", "Bedrooms",
        "Elec (kWh)", "Gas (therm)", "Gross Elec ($)", "Gross Gas ($)",
        "PV Credit ($)", "Net Cost ($)"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.style = styles["header"]

    # Data rows
    for row_idx, zone in enumerate(report.zone_results, 2):
        ws.cell(row=row_idx, column=1, value=zone.zone_name)
        ws.cell(row=row_idx, column=2, value="Dwelling" if zone.is_dwelling_unit else "Common")
        ws.cell(row=row_idx, column=3, value=zone.category.value if zone.category else "")
        ws.cell(row=row_idx, column=4, value=zone.area_sqft).style = styles["number"]
        ws.cell(row=row_idx, column=5, value=zone.num_bedrooms if zone.is_dwelling_unit else "")
        ws.cell(row=row_idx, column=6, value=zone.annual_elec_kwh).style = styles["number"]
        ws.cell(row=row_idx, column=7, value=zone.annual_gas_therm).style = styles["number"]
        ws.cell(row=row_idx, column=8, value=zone.gross_elec_cost).style = styles["currency"]
        ws.cell(row=row_idx, column=9, value=zone.gross_gas_cost).style = styles["currency"]
        ws.cell(row=row_idx, column=10, value=zone.pv_credit).style = styles["currency"]
        ws.cell(row=row_idx, column=11, value=zone.net_total_cost).style = styles["currency"]

        # Alternate row coloring
        if row_idx % 2 == 0:
            for col in range(1, 12):
                ws.cell(row=row_idx, column=col).fill = PatternFill(
                    start_color=COLORS['light_gray'],
                    end_color=COLORS['light_gray'],
                    fill_type="solid"
                )

    # Totals row
    total_row = len(report.zone_results) + 2
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=report.total_area_sqft).style = styles["number"]
    ws.cell(row=total_row, column=6, value=report.total_elec_kwh).style = styles["number"]
    ws.cell(row=total_row, column=7, value=report.total_gas_therm).style = styles["number"]
    ws.cell(row=total_row, column=11, value=report.total_annual_cost).style = styles["currency"]

    # Bold totals
    for col in range(1, 12):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    # Adjust column widths
    col_widths = [30, 10, 12, 12, 10, 12, 10, 12, 12, 12, 12]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Freeze header row
    ws.freeze_panes = 'A2'


def _create_dwelling_units_sheet(
    wb: "Workbook",
    report: "ZoneLccaReport",
    styles: Dict[str, Any],
    include_charts: bool = True
):
    """Create dwelling units breakdown sheet."""
    ws = wb.create_sheet("Dwelling Units")

    if not report.bedroom_summaries:
        ws['A1'] = "No dwelling unit data available."
        return

    # Headers
    headers = [
        "Bedrooms", "Unit Count", "Total Area", "Avg Area",
        "Avg Elec (kWh)", "Avg Gas (therm)", "Avg Annual Cost",
        "Utility Allow.", "PV (kWdc)", "Battery (kWh)"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.style = styles["header"]

    # Data rows
    for row_idx, bs in enumerate(report.bedroom_summaries, 2):
        ws.cell(row=row_idx, column=1, value=f"{bs.num_bedrooms}-BR")
        ws.cell(row=row_idx, column=2, value=bs.unit_count).style = styles["number"]
        ws.cell(row=row_idx, column=3, value=bs.total_area_sqft).style = styles["number"]
        ws.cell(row=row_idx, column=4, value=bs.avg_area_sqft).style = styles["number"]
        ws.cell(row=row_idx, column=5, value=bs.avg_elec_kwh).style = styles["number"]
        ws.cell(row=row_idx, column=6, value=bs.avg_gas_therm).style = styles["decimal"]
        ws.cell(row=row_idx, column=7, value=bs.avg_annual_cost).style = styles["currency"]
        ws.cell(row=row_idx, column=8, value=bs.utility_allowance_monthly).style = styles["currency"]
        ws.cell(row=row_idx, column=9, value=bs.pv_per_unit_kwdc).style = styles["decimal"]
        ws.cell(row=row_idx, column=10, value=bs.battery_per_unit_kwh).style = styles["decimal"]

    # Adjust column widths
    col_widths = [10, 12, 12, 10, 14, 14, 14, 14, 12, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Add bar chart if requested
    if include_charts and len(report.bedroom_summaries) > 0:
        chart = BarChart()
        chart.title = "Average Annual Cost by Bedroom Count"
        chart.y_axis.title = "Annual Cost ($)"
        chart.x_axis.title = "Unit Type"

        data_end_row = len(report.bedroom_summaries) + 1
        data = Reference(ws, min_col=7, min_row=1, max_row=data_end_row, max_col=7)
        cats = Reference(ws, min_col=1, min_row=2, max_row=data_end_row)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        chart.width = 12
        chart.height = 8

        ws.add_chart(chart, "L2")


def _create_common_areas_sheet(
    wb: "Workbook",
    report: "ZoneLccaReport",
    styles: Dict[str, Any],
    include_charts: bool = True
):
    """Create common areas breakdown sheet."""
    ws = wb.create_sheet("Common Areas")

    if not report.common_area_summaries:
        ws['A1'] = "No common area data available."
        return

    # Headers
    headers = [
        "Category", "Zone Count", "Area (sqft)", "Elec (kWh)",
        "Gas (therm)", "Annual Cost", "EUI (kBtu/sqft)", "$/sqft"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.style = styles["header"]

    # Data rows
    for row_idx, cas in enumerate(report.common_area_summaries, 2):
        ws.cell(row=row_idx, column=1, value=cas.category.value.title())
        ws.cell(row=row_idx, column=2, value=cas.zone_count).style = styles["number"]
        ws.cell(row=row_idx, column=3, value=cas.total_area_sqft).style = styles["number"]
        ws.cell(row=row_idx, column=4, value=cas.total_elec_kwh).style = styles["number"]
        ws.cell(row=row_idx, column=5, value=cas.total_gas_therm).style = styles["decimal"]
        ws.cell(row=row_idx, column=6, value=cas.total_annual_cost).style = styles["currency"]
        ws.cell(row=row_idx, column=7, value=cas.eui_kbtu_sqft).style = styles["decimal"]
        ws.cell(row=row_idx, column=8, value=cas.cost_per_sqft).style = styles["currency"]

    # Adjust column widths
    col_widths = [15, 12, 12, 12, 12, 14, 14, 10]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Add pie chart for cost breakdown
    if include_charts and len(report.common_area_summaries) > 1:
        chart = PieChart()
        chart.title = "Common Area Costs by Category"

        data_end_row = len(report.common_area_summaries) + 1
        data = Reference(ws, min_col=6, min_row=1, max_row=data_end_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=data_end_row)

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 12
        chart.height = 10

        # Add data labels
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = False

        ws.add_chart(chart, "J2")


def _create_cuac_sheet(
    wb: "Workbook",
    report: "ZoneLccaReport",
    styles: Dict[str, Any]
):
    """Create CUAC utility allowance sheet."""
    ws = wb.create_sheet("CUAC Allowances")

    if not report.has_cuac:
        ws['A1'] = "No CUAC data available."
        return

    cuac = report.cuac_summary
    row = 1

    # Title
    ws.merge_cells(f'A{row}:G{row}')
    ws[f'A{row}'] = "CUAC UTILITY ALLOWANCE SCHEDULE"
    ws[f'A{row}'].style = styles["title"]
    row += 2

    # Info
    ws[f'A{row}'] = "Utility:"
    ws[f'B{row}'] = cuac.utility_name
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    ws[f'A{row}'] = "Tariff:"
    ws[f'B{row}'] = cuac.tariff_name
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    ws[f'A{row}'] = "Report Date:"
    ws[f'B{row}'] = cuac.report_date
    ws[f'A{row}'].font = Font(bold=True)
    row += 1

    if cuac.total_pv_kwdc > 0:
        ws[f'A{row}'] = "PV System:"
        ws[f'B{row}'] = f"{cuac.total_pv_kwdc:.1f} kWdc"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

        ws[f'A{row}'] = "PV Billing:"
        ws[f'B{row}'] = cuac.pv_billing_option
        ws[f'A{row}'].font = Font(bold=True)
        row += 1

    row += 1

    # Allowance table headers
    headers = [
        "Bedrooms", "Electric ($)", "Gas ($)", "Gross Total ($)",
        "PV Credit ($)", "Batt Credit ($)", "Net Allowance ($)"
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.style = styles["header"]

    row += 1

    # Allowance rows
    for allowance_row in cuac.get_allowance_table():
        ws.cell(row=row, column=1, value=f"{allowance_row['bedrooms']}-BR")
        ws.cell(row=row, column=2, value=allowance_row['gross_elec']).style = styles["currency"]
        ws.cell(row=row, column=3, value=allowance_row['gross_gas']).style = styles["currency"]
        ws.cell(row=row, column=4, value=allowance_row['gross_total']).style = styles["currency"]
        ws.cell(row=row, column=5, value=allowance_row['pv_credit']).style = styles["currency"]
        ws.cell(row=row, column=6, value=allowance_row['battery_credit']).style = styles["currency"]
        ws.cell(row=row, column=7, value=allowance_row['net_allowance']).style = styles["currency"]

        # Highlight net allowance column
        ws.cell(row=row, column=7).font = Font(bold=True)

        row += 1

    # Adjust column widths
    col_widths = [12, 12, 12, 14, 12, 14, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
