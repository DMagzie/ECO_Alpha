"""
Excel Dashboard Export for LCCA.

Provides:
- LCCA results export to formatted Excel workbook
- Cash flow tables with charts
- Energy cost breakdown sheets
- Summary dashboard
"""

from __future__ import annotations
from dataclasses import dataclass
from typing import List, Optional, Dict, Any
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, NamedStyle
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.label import DataLabelList
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .calculators import LccaResults
from .econ1 import Econ1Report
from .tariffs import TouCostBreakdown


@dataclass
class ExcelExportOptions:
    """Options for Excel export."""
    include_cash_flow_chart: bool = True
    include_energy_breakdown: bool = True
    include_tou_breakdown: bool = True
    include_cost_summary: bool = True
    company_name: str = ""
    prepared_by: str = ""


def _check_openpyxl():
    """Check if openpyxl is available."""
    if not OPENPYXL_AVAILABLE:
        raise ImportError(
            "openpyxl is required for Excel export. "
            "Install with: pip install openpyxl"
        )


def _create_styles(wb: "Workbook") -> Dict[str, "NamedStyle"]:
    """Create named styles for the workbook."""
    styles = {}

    # Title style
    title_style = NamedStyle(name="title_style")
    title_style.font = Font(bold=True, size=16)
    title_style.alignment = Alignment(horizontal="center")
    wb.add_named_style(title_style)
    styles["title"] = title_style

    # Header style
    header_style = NamedStyle(name="header_style")
    header_style.font = Font(bold=True, size=11, color="FFFFFF")
    header_style.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_style.alignment = Alignment(horizontal="center", vertical="center")
    header_style.border = Border(
        bottom=Side(style="thin", color="000000")
    )
    wb.add_named_style(header_style)
    styles["header"] = header_style

    # Currency style
    currency_style = NamedStyle(name="currency_style")
    currency_style.number_format = '"$"#,##0'
    currency_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(currency_style)
    styles["currency"] = currency_style

    # Percent style
    percent_style = NamedStyle(name="percent_style")
    percent_style.number_format = "0.0%"
    percent_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(percent_style)
    styles["percent"] = percent_style

    # Number style
    number_style = NamedStyle(name="number_style")
    number_style.number_format = "#,##0"
    number_style.alignment = Alignment(horizontal="right")
    wb.add_named_style(number_style)
    styles["number"] = number_style

    # Section header style
    section_style = NamedStyle(name="section_style")
    section_style.font = Font(bold=True, size=12)
    section_style.fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    wb.add_named_style(section_style)
    styles["section"] = section_style

    return styles


def export_lcca_to_excel(
    results: LccaResults,
    file_path: str,
    project_name: str = "LCCA Analysis",
    econ1: Optional[Econ1Report] = None,
    tou_breakdown: Optional[TouCostBreakdown] = None,
    options: Optional[ExcelExportOptions] = None
) -> str:
    """
    Export LCCA results to formatted Excel workbook.

    Args:
        results: LccaResults from run_lcca()
        file_path: Output file path (.xlsx)
        project_name: Project name for header
        econ1: Optional ECON-1 report data
        tou_breakdown: Optional TOU cost breakdown
        options: Export options

    Returns:
        Path to created file
    """
    _check_openpyxl()

    if options is None:
        options = ExcelExportOptions()

    wb = Workbook()
    styles = _create_styles(wb)

    # Create sheets
    _create_summary_sheet(wb, results, project_name, options)
    _create_cash_flow_sheet(wb, results, options)

    if econ1 and options.include_energy_breakdown:
        _create_energy_sheet(wb, econ1)

    if tou_breakdown and options.include_tou_breakdown:
        _create_tou_sheet(wb, tou_breakdown)

    # Remove default sheet if others exist
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    # Save
    wb.save(file_path)
    return file_path


def _create_summary_sheet(
    wb: "Workbook",
    results: LccaResults,
    project_name: str,
    options: ExcelExportOptions
):
    """Create summary dashboard sheet."""
    ws = wb.create_sheet("Summary", 0)

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = "LIFE CYCLE COST ANALYSIS"
    ws["A1"].style = "title_style"

    ws.merge_cells("A2:F2")
    ws["A2"] = project_name
    ws["A2"].font = Font(size=14, italic=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    # Prepared info
    row = 4
    if options.company_name:
        ws[f"A{row}"] = "Prepared for:"
        ws[f"B{row}"] = options.company_name
        row += 1
    if options.prepared_by:
        ws[f"A{row}"] = "Prepared by:"
        ws[f"B{row}"] = options.prepared_by
        row += 1

    # Investment Summary
    row += 2
    ws[f"A{row}"] = "INVESTMENT SUMMARY"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:C{row}")

    row += 1
    data = [
        ("Initial Investment", results.initial_investment, "currency"),
        ("Total Incentives (PV)", results.total_incentives, "currency"),
        ("Net Investment", results.net_investment, "currency"),
    ]

    for label, value, fmt in data:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = f"{fmt}_style"
        row += 1

    # Annual Costs
    row += 1
    ws[f"A{row}"] = "ANNUAL COSTS (Year 1)"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:C{row}")

    row += 1
    data = [
        ("Baseline Annual Cost", results.baseline_annual_cost, "currency"),
        ("Proposed Annual Cost", results.proposed_annual_cost, "currency"),
        ("Annual Savings", results.annual_savings, "currency"),
    ]

    for label, value, fmt in data:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = f"{fmt}_style"
        row += 1

    # Financial Metrics
    row += 1
    ws[f"A{row}"] = "FINANCIAL METRICS"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:C{row}")

    row += 1
    ws[f"A{row}"] = "Net Present Value (NPV)"
    ws[f"C{row}"] = results.npv
    ws[f"C{row}"].style = "currency_style"
    row += 1

    ws[f"A{row}"] = "Internal Rate of Return (IRR)"
    if results.irr is not None:
        ws[f"C{row}"] = results.irr
        ws[f"C{row}"].style = "percent_style"
    else:
        ws[f"C{row}"] = "N/A"
    row += 1

    ws[f"A{row}"] = "Simple Payback (years)"
    if results.simple_payback_years is not None:
        ws[f"C{row}"] = results.simple_payback_years
        ws[f"C{row}"].number_format = "0.0"
    else:
        ws[f"C{row}"] = "N/A"
    row += 1

    ws[f"A{row}"] = "Discounted Payback (years)"
    if results.discounted_payback_years is not None:
        ws[f"C{row}"] = results.discounted_payback_years
        ws[f"C{row}"].number_format = "0.0"
    else:
        ws[f"C{row}"] = "N/A"
    row += 1

    ws[f"A{row}"] = "Savings-to-Investment Ratio (SIR)"
    if results.sir is not None:
        ws[f"C{row}"] = results.sir
        ws[f"C{row}"].number_format = "0.00"
    else:
        ws[f"C{row}"] = "N/A"
    row += 1

    # Lifecycle Costs
    row += 1
    ws[f"A{row}"] = "LIFECYCLE COSTS (Present Value)"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:C{row}")

    row += 1
    data = [
        ("Baseline Lifecycle Cost", results.baseline_lifecycle_cost, "currency"),
        ("Proposed Lifecycle Cost", results.proposed_lifecycle_cost, "currency"),
        ("Lifecycle Savings", results.lifecycle_savings, "currency"),
    ]

    for label, value, fmt in data:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = f"{fmt}_style"
        row += 1

    # Analysis Parameters
    row += 1
    ws[f"A{row}"] = "ANALYSIS PARAMETERS"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:C{row}")

    row += 1
    ws[f"A{row}"] = "Analysis Period"
    ws[f"C{row}"] = f"{results.analysis_years} years"
    row += 1
    ws[f"A{row}"] = "Real Discount Rate"
    ws[f"C{row}"] = results.discount_rate
    ws[f"C{row}"].style = "percent_style"

    # Column widths
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 18


def _create_cash_flow_sheet(
    wb: "Workbook",
    results: LccaResults,
    options: ExcelExportOptions
):
    """Create cash flow table and chart."""
    ws = wb.create_sheet("Cash Flow")

    # Title
    ws["A1"] = "Annual Cash Flow Analysis"
    ws["A1"].style = "title_style"
    ws.merge_cells("A1:E1")

    # Headers
    headers = ["Year", "Cash Flow", "Cumulative", "Category"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.style = "header_style"

    # Data
    row = 4
    for i, cf in enumerate(results.cash_flows):
        ws.cell(row=row, column=1, value=cf.year)
        ws.cell(row=row, column=2, value=cf.amount).style = "currency_style"
        if i < len(results.cumulative_cash_flows):
            ws.cell(row=row, column=3, value=results.cumulative_cash_flows[i]).style = "currency_style"
        ws.cell(row=row, column=4, value=cf.category.title())
        row += 1

    # Add chart if requested
    if options.include_cash_flow_chart and len(results.cash_flows) > 1:
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "Cash Flow by Year"
        chart.y_axis.title = "Dollars ($)"
        chart.x_axis.title = "Year"

        data = Reference(ws, min_col=2, min_row=3, max_row=row-1, max_col=2)
        cats = Reference(ws, min_col=1, min_row=4, max_row=row-1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        chart.width = 15
        chart.height = 10

        ws.add_chart(chart, "F3")

        # Cumulative line chart
        line_chart = LineChart()
        line_chart.title = "Cumulative Cash Flow"
        line_chart.y_axis.title = "Dollars ($)"
        line_chart.x_axis.title = "Year"
        line_chart.style = 10

        data2 = Reference(ws, min_col=3, min_row=3, max_row=row-1, max_col=3)
        line_chart.add_data(data2, titles_from_data=True)
        line_chart.set_categories(cats)
        line_chart.width = 15
        line_chart.height = 10

        ws.add_chart(line_chart, "F20")

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 12


def _create_energy_sheet(wb: "Workbook", econ1: Econ1Report):
    """Create energy breakdown sheet from ECON-1 data."""
    ws = wb.create_sheet("Energy Breakdown")

    # Title
    ws["A1"] = "Energy Cost Breakdown"
    ws["A1"].style = "title_style"
    ws.merge_cells("A1:D1")

    # Project info
    row = 3
    ws[f"A{row}"] = "PROJECT INFORMATION"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    info = [
        ("Project Name", econ1.project_name),
        ("Building Type", econ1.building_type),
        ("Conditioned Area (SF)", f"{econ1.conditioned_area_sf:,.0f}"),
    ]
    for label, value in info:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        row += 1

    # Proposed Energy
    row += 1
    ws[f"A{row}"] = "PROPOSED BUILDING - ANNUAL ENERGY"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    proposed = econ1.proposed
    energy_data = [
        ("Electricity (kWh)", proposed.elec_kwh_gross, "number"),
        ("Natural Gas (therms)", proposed.gas_therm, "number"),
        ("Peak Demand (kW)", proposed.peak_demand_kw, "number"),
    ]
    if proposed.generation and proposed.generation.pv_generation_kwh > 0:
        energy_data.append(("PV Generation (kWh)", proposed.generation.pv_generation_kwh, "number"))

    for label, value, fmt in energy_data:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = f"{fmt}_style"
        row += 1

    # Proposed Costs
    row += 1
    ws[f"A{row}"] = "PROPOSED BUILDING - ANNUAL COSTS"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    cost_data = [
        ("Electricity Cost", proposed.elec_cost_gross, "currency"),
        ("Natural Gas Cost", proposed.gas_cost, "currency"),
        ("Demand Charges", proposed.demand_cost, "currency"),
        ("Gross Total", proposed.gross_cost, "currency"),
    ]
    if proposed.pv_savings > 0:
        cost_data.append(("PV Credit", -proposed.pv_savings, "currency"))
    cost_data.append(("Net Annual Cost", proposed.net_cost, "currency"))

    for label, value, fmt in cost_data:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = f"{fmt}_style"
        row += 1

    # Baseline if available
    if econ1.baseline:
        row += 1
        ws[f"A{row}"] = "BASELINE BUILDING - ANNUAL COSTS"
        ws[f"A{row}"].style = "section_style"
        ws.merge_cells(f"A{row}:D{row}")

        row += 1
        baseline = econ1.baseline
        baseline_data = [
            ("Electricity Cost", baseline.elec_cost_gross, "currency"),
            ("Natural Gas Cost", baseline.gas_cost, "currency"),
            ("Demand Charges", baseline.demand_cost, "currency"),
            ("Net Annual Cost", baseline.net_cost, "currency"),
        ]

        for label, value, fmt in baseline_data:
            ws[f"A{row}"] = label
            ws[f"C{row}"] = value
            ws[f"C{row}"].style = f"{fmt}_style"
            row += 1

        # Savings
        row += 1
        ws[f"A{row}"] = "COST COMPARISON"
        ws[f"A{row}"].style = "section_style"
        ws.merge_cells(f"A{row}:D{row}")

        row += 1
        ws[f"A{row}"] = "Annual Savings"
        ws[f"C{row}"] = econ1.annual_savings
        ws[f"C{row}"].style = "currency_style"
        row += 1
        ws[f"A{row}"] = "Percent Savings"
        ws[f"C{row}"] = econ1.percent_savings / 100 if econ1.percent_savings else 0
        ws[f"C{row}"].style = "percent_style"

    # Column widths
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 18


def _create_tou_sheet(wb: "Workbook", tou: TouCostBreakdown):
    """Create TOU breakdown sheet."""
    ws = wb.create_sheet("TOU Breakdown")

    # Title
    ws["A1"] = "Time-of-Use Energy Cost Breakdown"
    ws["A1"].style = "title_style"
    ws.merge_cells("A1:D1")

    # Summer
    row = 3
    ws[f"A{row}"] = "SUMMER ENERGY CHARGES"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    headers = ["Period", "kWh", "Cost"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=h).style = "header_style"

    row += 1
    summer_data = [
        ("On-Peak", tou.summer_on_peak_kwh, tou.summer_on_peak_cost),
        ("Mid-Peak", tou.summer_mid_peak_kwh, tou.summer_mid_peak_cost),
        ("Off-Peak", tou.summer_off_peak_kwh, tou.summer_off_peak_cost),
    ]
    for period, kwh, cost in summer_data:
        ws.cell(row=row, column=1, value=period)
        ws.cell(row=row, column=2, value=kwh).style = "number_style"
        ws.cell(row=row, column=3, value=cost).style = "currency_style"
        row += 1

    # Winter
    row += 1
    ws[f"A{row}"] = "WINTER ENERGY CHARGES"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    for col, h in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=h).style = "header_style"

    row += 1
    winter_data = [
        ("On-Peak", tou.winter_on_peak_kwh, tou.winter_on_peak_cost),
        ("Mid-Peak", tou.winter_mid_peak_kwh, tou.winter_mid_peak_cost),
        ("Off-Peak", tou.winter_off_peak_kwh, tou.winter_off_peak_cost),
    ]
    for period, kwh, cost in winter_data:
        ws.cell(row=row, column=1, value=period)
        ws.cell(row=row, column=2, value=kwh).style = "number_style"
        ws.cell(row=row, column=3, value=cost).style = "currency_style"
        row += 1

    # Demand
    row += 1
    ws[f"A{row}"] = "DEMAND CHARGES"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    demand_headers = ["Type", "kW", "Cost"]
    for col, h in enumerate(demand_headers, 1):
        ws.cell(row=row, column=col, value=h).style = "header_style"

    row += 1
    demand_data = [
        ("Facility", tou.facility_demand_kw, tou.facility_demand_cost),
        ("Summer On-Peak", tou.summer_on_peak_demand_kw, tou.summer_on_peak_demand_cost),
        ("Summer Mid-Peak", tou.summer_mid_peak_demand_kw, tou.summer_mid_peak_demand_cost),
    ]
    for dtype, kw, cost in demand_data:
        if cost > 0:
            ws.cell(row=row, column=1, value=dtype)
            ws.cell(row=row, column=2, value=kw).style = "number_style"
            ws.cell(row=row, column=3, value=cost).style = "currency_style"
            row += 1

    # Summary
    row += 1
    ws[f"A{row}"] = "SUMMARY"
    ws[f"A{row}"].style = "section_style"
    ws.merge_cells(f"A{row}:D{row}")

    row += 1
    summary = [
        ("Total Energy Cost", tou.total_energy_cost),
        ("Total Demand Cost", tou.total_demand_cost),
        ("Fixed Charges", tou.total_fixed_cost),
        ("ANNUAL TOTAL", tou.total_cost),
    ]
    for label, value in summary:
        ws[f"A{row}"] = label
        ws[f"C{row}"] = value
        ws[f"C{row}"].style = "currency_style"
        if label == "ANNUAL TOTAL":
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"C{row}"].font = Font(bold=True)
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15


def export_econ1_to_excel(
    econ1: Econ1Report,
    file_path: str,
    options: Optional[ExcelExportOptions] = None
) -> str:
    """
    Export ECON-1 report to Excel workbook.

    Args:
        econ1: Econ1Report object
        file_path: Output file path (.xlsx)
        options: Export options

    Returns:
        Path to created file
    """
    _check_openpyxl()

    if options is None:
        options = ExcelExportOptions()

    wb = Workbook()
    styles = _create_styles(wb)

    # Main sheet
    ws = wb.active
    ws.title = "ECON-1"

    # Title
    ws["A1"] = "ECON-1 ENERGY COST ANALYSIS"
    ws["A1"].style = "title_style"
    ws.merge_cells("A1:D1")

    _create_energy_sheet(wb, econ1)

    # Remove default if we created others
    if len(wb.sheetnames) > 1:
        del wb["ECON-1"]

    wb.save(file_path)
    return file_path


def export_comparison_to_excel(
    scenarios: List[Dict[str, Any]],
    file_path: str,
    title: str = "Scenario Comparison"
) -> str:
    """
    Export multiple scenarios to comparison Excel workbook.

    Args:
        scenarios: List of scenario dicts with 'name', 'results' (LccaResults)
        file_path: Output file path
        title: Report title

    Returns:
        Path to created file
    """
    _check_openpyxl()

    wb = Workbook()
    styles = _create_styles(wb)
    ws = wb.active
    ws.title = "Comparison"

    # Title
    ws["A1"] = title
    ws["A1"].style = "title_style"
    ws.merge_cells(f"A1:{get_column_letter(len(scenarios) + 1)}1")

    # Headers
    row = 3
    ws.cell(row=row, column=1, value="Metric").style = "header_style"
    for col, scenario in enumerate(scenarios, 2):
        ws.cell(row=row, column=col, value=scenario["name"]).style = "header_style"

    # Metrics
    metrics = [
        ("Initial Investment", "initial_investment", "currency"),
        ("Net Investment", "net_investment", "currency"),
        ("Annual Savings", "annual_savings", "currency"),
        ("NPV", "npv", "currency"),
        ("IRR", "irr", "percent"),
        ("Simple Payback (years)", "simple_payback_years", "number"),
        ("SIR", "sir", "number"),
        ("Lifecycle Savings", "lifecycle_savings", "currency"),
    ]

    row = 4
    for label, attr, fmt in metrics:
        ws.cell(row=row, column=1, value=label)
        for col, scenario in enumerate(scenarios, 2):
            results = scenario["results"]
            value = getattr(results, attr, None)
            cell = ws.cell(row=row, column=col, value=value)
            if fmt == "currency":
                cell.style = "currency_style"
            elif fmt == "percent" and value is not None:
                cell.style = "percent_style"
            elif fmt == "number" and value is not None:
                cell.number_format = "0.00"
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 25
    for col in range(2, len(scenarios) + 2):
        ws.column_dimensions[get_column_letter(col)].width = 18

    wb.save(file_path)
    return file_path
